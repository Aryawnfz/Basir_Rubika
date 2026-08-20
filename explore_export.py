"""
خروجی اکسل برای آمارِ کاوشِ کانال‌های روبیکا (نتیجهٔ analyze_channel).
"""
from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

_HEADER_FILL = PatternFill("solid", fgColor="1a1a3e")
_HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=12)
_LABEL_FONT = Font(bold=True, color="1a1a3e", name="Calibri", size=11)
_ROW_FILL_ODD = PatternFill("solid", fgColor="F0F0F8")
_ROW_FILL_EVEN = PatternFill("solid", fgColor="FFFFFF")
_SIDE = Side(border_style="thin", color="C0C0D0")
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)


def generate_channel_stats_excel(stats: dict) -> BytesIO:
    wb = openpyxl.Workbook()

    # ── شیت خلاصهٔ آمار ─────────────────────────────────────────────
    ws = wb.active
    ws.title = "خلاصه آمار کانال"
    ws.sheet_view.rightToLeft = True

    ws.merge_cells("A1:B1")
    tcell = ws.cell(row=1, column=1, value=f"آمار کانال: @{stats.get('username','')}")
    tcell.font = _HEADER_FONT
    tcell.fill = _HEADER_FILL
    tcell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24

    rows = [
        ("تعداد پست‌های بررسی‌شده", stats.get("posts_analyzed", 0)),
        ("مجموع بازدیدها", stats.get("total_views", 0)),
        ("مجموع ری‌اکشن‌ها", stats.get("total_reactions", 0)),
        ("میانگین بازدید هر پست", stats.get("avg_views", 0)),
        ("میانگین ری‌اکشن هر پست", stats.get("avg_reactions", 0)),
        ("نرخ تعامل (٪)", stats.get("engagement_rate", 0)),
    ]
    for i, (label, value) in enumerate(rows, start=2):
        lc = ws.cell(row=i, column=1, value=label)
        vc = ws.cell(row=i, column=2, value=value)
        lc.font = _LABEL_FONT
        lc.alignment = Alignment(horizontal="right", vertical="center")
        vc.alignment = Alignment(horizontal="center", vertical="center")
        for c in (lc, vc):
            c.border = _BORDER
            c.fill = _ROW_FILL_ODD if i % 2 == 0 else _ROW_FILL_EVEN

    # ── شیت پست‌ها ──────────────────────────────────────────────────
    ws2 = wb.create_sheet("پست‌ها")
    ws2.sheet_view.rightToLeft = True
    headers = ["#", "محتوای پست", "تاریخ پست", "بازدید",
               "مجموع ری‌اکشن", "ری‌اکشن‌ها", "لینک پست"]
    widths = [5, 60, 18, 12, 14, 30, 42]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
        ws2.column_dimensions[c.column_letter].width = w
    ws2.row_dimensions[1].height = 26

    for r, post in enumerate(stats.get("posts", []), start=2):
        reactions = post.get("reactions", []) or []
        reactions_str = " | ".join(
            f"{rx.get('emoji','')} {rx.get('count','')}" for rx in reactions
        )
        values = [
            r - 1,
            post.get("content", ""),
            post.get("datetime_full", ""),
            post.get("views", ""),
            post.get("total_reactions", 0),
            reactions_str,
            post.get("link", ""),
        ]
        fill = _ROW_FILL_ODD if r % 2 == 0 else _ROW_FILL_EVEN
        for col, val in enumerate(values, start=1):
            c = ws2.cell(row=r, column=col, value=val)
            c.fill = fill
            c.border = _BORDER
            c.alignment = Alignment(
                horizontal="right" if col == 2 else "center",
                vertical="center", wrap_text=True,
            )
        ws2.row_dimensions[r].height = 38
    ws2.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_all_channels_excel(records: list[dict]) -> BytesIO:
    """
    خروجی کلی: آمارِ همهٔ کانال‌های بررسی‌شده در یک شیتِ خلاصه (یک ردیف
    به‌ازای هر کانال، همراه تعداد اعضا) + یک شیتِ پست‌ها که پست‌های همهٔ
    کانال‌ها را کنار هم دارد.
    """
    wb = openpyxl.Workbook()

    # ── شیت خلاصهٔ همهٔ کانال‌ها ─────────────────────────────────────
    ws = wb.active
    ws.title = "خلاصه کانال‌ها"
    ws.sheet_view.rightToLeft = True
    headers = ["#", "کانال", "آیدی", "تعداد اعضا", "پست بررسی‌شده",
               "مجموع بازدید", "مجموع ری‌اکشن", "میانگین بازدید",
               "میانگین ری‌اکشن", "نرخ تعامل (٪)"]
    widths = [5, 26, 20, 16, 14, 14, 15, 14, 15, 14]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30

    tot_views = tot_reactions = tot_posts = 0
    for r, rec in enumerate(records, start=2):
        res = rec.get("result", {}) or {}
        members = rec.get("members", "") or (
            f"{rec.get('members_count', 0)}" if rec.get("members_count") else ""
        )
        tot_views += res.get("total_views", 0)
        tot_reactions += res.get("total_reactions", 0)
        tot_posts += res.get("posts_analyzed", 0)
        values = [
            r - 1,
            rec.get("title", "") or rec.get("username", ""),
            f"@{rec.get('username','')}",
            members,
            res.get("posts_analyzed", 0),
            res.get("total_views", 0),
            res.get("total_reactions", 0),
            res.get("avg_views", 0),
            res.get("avg_reactions", 0),
            res.get("engagement_rate", 0),
        ]
        fill = _ROW_FILL_ODD if r % 2 == 0 else _ROW_FILL_EVEN
        for col, val in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=val)
            c.fill = fill
            c.border = _BORDER
            c.alignment = Alignment(
                horizontal="right" if col in (2, 3) else "center",
                vertical="center", wrap_text=True,
            )

    trow = len(records) + 2
    total_engagement = round((tot_reactions / tot_views) * 100, 2) if tot_views else 0
    avg_views = round(tot_views / tot_posts, 1) if tot_posts else 0
    avg_reactions = round(tot_reactions / tot_posts, 1) if tot_posts else 0
    totals = ["", "مجموع", "", "", tot_posts, tot_views, tot_reactions,
              avg_views, avg_reactions, total_engagement]
    for col, val in enumerate(totals, start=1):
        c = ws.cell(row=trow, column=col, value=val)
        c.font = _LABEL_FONT
        c.fill = PatternFill("solid", fgColor="E4E4F4")
        c.border = _BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    # ── شیت پست‌ها (همهٔ کانال‌ها) ──────────────────────────────────
    ws2 = wb.create_sheet("پست‌ها")
    ws2.sheet_view.rightToLeft = True
    headers2 = ["#", "کانال", "محتوای پست", "تاریخ پست", "بازدید",
                "مجموع ری‌اکشن", "ری‌اکشن‌ها", "لینک پست"]
    widths2 = [5, 22, 55, 18, 12, 14, 28, 40]
    for col, (h, w) in enumerate(zip(headers2, widths2), start=1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
        ws2.column_dimensions[c.column_letter].width = w
    ws2.row_dimensions[1].height = 26

    r = 2
    for rec in records:
        res = rec.get("result", {}) or {}
        ch_title = rec.get("title", "") or rec.get("username", "")
        for post in res.get("posts", []):
            reactions = post.get("reactions", []) or []
            reactions_str = " | ".join(
                f"{rx.get('emoji','')} {rx.get('count','')}" for rx in reactions
            )
            values = [
                r - 1, ch_title, post.get("content", ""),
                post.get("datetime_full", ""), post.get("views", ""),
                post.get("total_reactions", 0), reactions_str, post.get("link", ""),
            ]
            fill = _ROW_FILL_ODD if r % 2 == 0 else _ROW_FILL_EVEN
            for col, val in enumerate(values, start=1):
                c = ws2.cell(row=r, column=col, value=val)
                c.fill = fill
                c.border = _BORDER
                c.alignment = Alignment(
                    horizontal="right" if col in (2, 3) else "center",
                    vertical="center", wrap_text=True,
                )
            ws2.row_dimensions[r].height = 38
            r += 1
    ws2.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
