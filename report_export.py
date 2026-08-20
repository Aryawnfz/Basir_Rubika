"""
خروجی اکسل «گزارش فعالیت‌ها» — دقیقاً همان چیزی که در صفحهٔ /reports
برای بازهٔ انتخاب‌شده (روز/هفته/ماه/همه) نمایش داده می‌شود.

سه شیت:
  ۱. خلاصه گزارش      — بازه، زمانِ تولید، کل جستجوها/نتایج، عبارت‌های یکتا
  ۲. پرتکرارترین جستجوها — یک ردیف برای هر عبارت (تعداد + آخرین جستجو)
  ۳. سهم اکانت‌ها      — تعداد و درصدِ پیام‌های پیداشده به‌ازای هر اکانت
"""
from datetime import datetime
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from rubika_search import _gregorian_to_jalali

_HEADER_FILL = PatternFill("solid", fgColor="1a1a3e")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_TITLE_FONT = Font(bold=True, color="1a1a3e", size=14)
_LABEL_FONT = Font(bold=True, color="333333", size=11)
_LABEL_FILL = PatternFill("solid", fgColor="EEF0FA")
_ROW_ODD = PatternFill("solid", fgColor="F0F0F8")
_ROW_EVEN = PatternFill("solid", fgColor="FFFFFF")
_SIDE = Side(border_style="thin", color="C0C0D0")
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)

_PERIOD_LABELS = {
    "day": "روز گذشته",
    "week": "هفتهٔ گذشته",
    "month": "ماه گذشته",
    "": "همهٔ زمان‌ها",
}


def _now_jalali_full() -> str:
    """«۱۴۰۵/۰۳/۲۷ ۰۸:۵۱ (۲۰۲۶-۰۶-۱۷ ۰۸:۵۱)» — برای هدرِ اکسل."""
    now = datetime.now()
    jy, jm, jd = _gregorian_to_jalali(now.year, now.month, now.day)
    return (
        f"{jy}/{jm:02d}/{jd:02d} {now.hour:02d}:{now.minute:02d} "
        f"({now.year}-{now.month:02d}-{now.day:02d} {now.hour:02d}:{now.minute:02d})"
    )


def _fmt_jalali_dt(ts: float) -> str:
    """یک Unix timestamp را به «۱۴۰۵/۰۵/۱۷ — ۰۲:۳۶» تبدیل می‌کند."""
    if not ts:
        return "—"
    dt = datetime.fromtimestamp(ts)
    jy, jm, jd = _gregorian_to_jalali(dt.year, dt.month, dt.day)
    return f"{jy}/{jm:02d}/{jd:02d} — {dt.hour:02d}:{dt.minute:02d}"


def _write_header(ws, headers, widths, row=1):
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 26


def _write_row(ws, row, values, right_cols=()):
    fill = _ROW_ODD if row % 2 == 0 else _ROW_EVEN
    for col, val in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=val)
        c.fill = fill
        c.border = _BORDER
        c.alignment = Alignment(
            horizontal="right" if col in right_cols else "center",
            vertical="center", wrap_text=True,
        )


def generate_report_excel(report: dict) -> BytesIO:
    period = report.get("period", "")
    period_label = _PERIOD_LABELS.get(period, _PERIOD_LABELS[""])
    queries = report.get("queries", [])
    accounts = report.get("account_slices", [])
    wb = openpyxl.Workbook()

    # ── شیت ۱: خلاصه گزارش ──────────────────────────────────────────
    ws = wb.active
    ws.title = "خلاصه گزارش"
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 30

    ws["A1"] = "📊 گزارش فعالیت‌های بصیر"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:B1")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    rows = [
        ("بازهٔ گزارش", period_label),
        ("زمان تولید گزارش", _now_jalali_full()),
        ("کل جستجوها", report.get("total_searches", 0)),
        ("کل نتایج", report.get("total_results", 0)),
        ("عبارت‌های یکتا", len(queries)),
        ("کل پیام‌های اکانت‌ها", report.get("total_account_msgs", 0)),
    ]
    for i, (label, value) in enumerate(rows, start=2):
        lc = ws.cell(row=i, column=1, value=label)
        vc = ws.cell(row=i, column=2, value=value)
        lc.font = _LABEL_FONT
        lc.fill = _LABEL_FILL
        lc.alignment = Alignment(horizontal="right", vertical="center")
        vc.alignment = Alignment(horizontal="center", vertical="center")
        for c in (lc, vc):
            c.border = _BORDER

    # ── شیت ۲: پرتکرارترین جستجوها ─────────────────────────────────
    ws2 = wb.create_sheet("پرتکرارترین جستجوها")
    ws2.sheet_view.rightToLeft = True
    _write_header(ws2, ["ردیف", "عبارت", "تعداد جستجو", "آخرین جستجو"],
                  [8, 42, 14, 28])
    for i, (query_text, stats) in enumerate(queries, start=1):
        _write_row(
            ws2, i + 1,
            [i, query_text, stats.get("count", 0), _fmt_jalali_dt(stats.get("last_at", 0))],
            right_cols=(2,),
        )
    ws2.freeze_panes = "A2"

    # ── شیت ۳: سهم اکانت‌ها ────────────────────────────────────────
    ws3 = wb.create_sheet("سهم اکانت‌ها")
    ws3.sheet_view.rightToLeft = True
    _write_header(ws3, ["ردیف", "اکانت", "تعداد پیام", "سهم (٪)"],
                  [8, 32, 14, 12])
    for i, acc in enumerate(accounts, start=1):
        _write_row(
            ws3, i + 1,
            [i, acc.get("name", ""), acc.get("count", 0), round(acc.get("pct", 0), 1)],
            right_cols=(2,),
        )
    ws3.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
