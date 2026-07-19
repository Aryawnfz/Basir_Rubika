"""
خروجی اکسل «آمار کلی» برای نتایج جستجوی بصیر.

یک فایل اکسل با چهار شیت تولید می‌کند که از روی نتایج یک جستجو محاسبه می‌شوند:
  ۱. خلاصه آمار        — شاخص‌های کلی + توزیع ری‌اکشن‌ها
  ۲. پست‌ها            — یک ردیف برای هر نتیجه
  ۳. آمار کانال‌ها     — تجمیع به‌ازای هر کانال
  ۴. توزیع روزانه      — تجمیع به‌ازای هر روز

نکته: فوروارد، کامنت و نوع رسانه توسط اسکریپر فعلی روبیکا جمع‌آوری نمی‌شوند،
پس این ستون‌ها صفر/خالی می‌مانند.
"""
import re
from collections import defaultdict
from datetime import datetime
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from rubika_search import (
    _gregorian_to_jalali,
    _jalali_to_gregorian,
    _to_latin_digits,
)

PLATFORM = "روبیکا"

# دسته‌بندی ایموجی ری‌اکشن‌ها
_POSITIVE = {
    "👍", "❤", "😍", "🔥", "👏", "🎉", "😂", "🙏", "💯", "😀", "😁", "😊",
    "🥰", "😘", "👌", "💪", "✅", "⭐", "🌟", "🤩", "💖", "💚", "💙", "💛",
    "🧡", "💜", "🤍", "🥳", "😄", "😃", "😇", "🤗", "😎", "🫡", "👀",
}
_NEGATIVE = {
    "👎", "😡", "😠", "😢", "😭", "💔", "🤬", "😞", "😔", "👿", "💩", "🤮",
    "😤", "🙄", "😒", "😩", "😫", "😟", "🤢", "😨", "😰",
}


def _strip_vs(emoji: str) -> str:
    """حذف variation selector و کاراکترهای صفرعرض برای مقایسهٔ پایدار ایموجی."""
    return emoji.replace("\ufe0f", "").replace("\u200d", "").strip()


def _reaction_kind(emoji: str) -> str:
    e = _strip_vs(emoji)
    if e in _POSITIVE:
        return "مثبت"
    if e in _NEGATIVE:
        return "منفی"
    return "خنثی"


def _parse_num(value) -> int:
    """رشته‌ای مثل «۳۳.۲K» یا «40,896» یا «819 بازدید» را به عدد صحیح تبدیل می‌کند."""
    if isinstance(value, (int, float)):
        return int(value)
    s = _to_latin_digits(str(value or "")).replace(",", "").replace("،", "")
    m = re.search(r"(\d+(?:\.\d+)?)\s*([KMkm])?", s)
    if not m:
        return 0
    num = float(m.group(1))
    suffix = (m.group(2) or "").upper()
    if suffix == "K":
        num *= 1_000
    elif suffix == "M":
        num *= 1_000_000
    return int(round(num))


def _fmt_int(n) -> str:
    return f"{int(round(n)):,}"


def _fmt_float(n) -> str:
    return f"{n:,.1f}"


def _split_dt(result: dict) -> tuple[str, str]:
    """
    از یک نتیجه، تاریخ شمسی (YYYY/MM/DD) و زمان (HH:MM) را برمی‌گرداند.
    اول از datetime_full، در غیر این صورت از date استفاده می‌کند.
    """
    raw = _to_latin_digits((result.get("datetime_full") or "").strip())
    if not raw:
        raw = _to_latin_digits((result.get("date") or "").strip())
    if not raw or raw in ("—", "-"):
        return "", ""
    parts = raw.split()
    jdate = parts[0] if parts else ""
    jtime = parts[1] if len(parts) > 1 else ""
    return jdate, jtime


def _greg_of(jdate: str) -> str:
    """تاریخ شمسی YYYY/MM/DD را به رشتهٔ میلادی YYYY-MM-DD تبدیل می‌کند."""
    m = re.match(r"^(\d{3,4})/(\d{1,2})/(\d{1,2})$", jdate or "")
    if not m:
        return ""
    jy, jm, jd = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        gy, gm, gd = _jalali_to_gregorian(jy, jm, jd)
        return f"{gy}-{gm:02d}-{gd:02d}"
    except Exception:
        return ""


def _fmt_dt_cell(jdate: str, jtime: str) -> str:
    """مثل نمونه: «1405/03/27 08:51 (2026-06-17 08:51)»."""
    if not jdate:
        return "—"
    greg = _greg_of(jdate)
    jpart = f"{jdate} {jtime}".strip()
    if greg:
        gpart = f"{greg} {jtime}".strip()
        return f"{jpart} ({gpart})"
    return jpart


# ── محاسبهٔ آمار ─────────────────────────────────────────────────────────────
def _compute(results: list[dict]) -> dict:
    total_views = 0
    total_reactions = 0
    pos = neg = neu = 0
    total_forwards = 0
    total_comments = 0
    emoji_counts: dict[str, int] = defaultdict(int)
    per_channel: dict[str, dict] = {}
    per_day: dict[str, dict] = {}

    rows = []
    for r in results:
        channel = r.get("channel_name", "") or "—"
        views = _parse_num(r.get("views", ""))
        text = (r.get("full_content") or r.get("message") or "").strip()
        jdate, jtime = _split_dt(r)

        r_pos = r_neg = r_neu = 0
        for rx in r.get("reactions", []) or []:
            emoji = (rx.get("emoji") or "").strip()
            cnt = _parse_num(rx.get("count", "0"))
            if not emoji or cnt <= 0:
                continue
            emoji_counts[_strip_vs(emoji)] += cnt
            kind = _reaction_kind(emoji)
            if kind == "مثبت":
                r_pos += cnt
            elif kind == "منفی":
                r_neg += cnt
            else:
                r_neu += cnt

        r_total = r_pos + r_neg + r_neu
        if r_total == 0:
            # اگر لیست ری‌اکشن خالی بود ولی total_reactions مقدار داشت
            r_total = _parse_num(r.get("total_reactions", 0))

        total_views += views
        total_reactions += r_total
        pos += r_pos
        neg += r_neg
        neu += r_neu

        rows.append({
            "channel": channel,
            "jdate": jdate,
            "jtime": jtime,
            "text": text,
            "views": views,
            "reactions": r_total,
            "pos": r_pos,
            "neg": r_neg,
            "forwards": 0,
            "comments": 0,
            "media": "—",
            "link": r.get("message_link", ""),
        })

        ch = per_channel.setdefault(
            channel, {"posts": 0, "views": 0, "reactions": 0, "forwards": 0}
        )
        ch["posts"] += 1
        ch["views"] += views
        ch["reactions"] += r_total

        if jdate:
            d = per_day.setdefault(jdate, {"posts": 0, "views": 0, "reactions": 0})
            d["posts"] += 1
            d["views"] += views
            d["reactions"] += r_total

    n = len(results)
    engagement_base = total_reactions + total_forwards + total_comments
    engagement = (engagement_base / total_views * 100) if total_views else 0.0

    return {
        "rows": rows,
        "n": n,
        "total_views": total_views,
        "total_reactions": total_reactions,
        "pos": pos,
        "neg": neg,
        "neu": neu,
        "total_forwards": total_forwards,
        "total_comments": total_comments,
        "avg_views": (total_views / n) if n else 0.0,
        "avg_reactions": (total_reactions / n) if n else 0.0,
        "engagement": engagement,
        "emoji_counts": emoji_counts,
        "per_channel": per_channel,
        "per_day": per_day,
    }


# ── استایل ───────────────────────────────────────────────────────────────────
_HEADER_FILL = PatternFill("solid", fgColor="1a1a3e")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_TITLE_FONT = Font(bold=True, color="1a1a3e", size=14)
_LABEL_FONT = Font(bold=True, color="333333", size=11)
_LABEL_FILL = PatternFill("solid", fgColor="EEF0FA")
_ROW_ODD = PatternFill("solid", fgColor="F0F0F8")
_ROW_EVEN = PatternFill("solid", fgColor="FFFFFF")
_SIDE = Side(border_style="thin", color="C0C0D0")
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)


def _write_header(ws, headers, widths, row=1):
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 26


def _now_jalali() -> str:
    now = datetime.now()
    jy, jm, jd = _gregorian_to_jalali(now.year, now.month, now.day)
    return (
        f"{jy}/{jm:02d}/{jd:02d} {now.hour:02d}:{now.minute:02d} "
        f"({now.year}-{now.month:02d}-{now.day:02d} {now.hour:02d}:{now.minute:02d})"
    )


def generate_stats_excel(results: list[dict], query: str = "") -> BytesIO:
    s = _compute(results)
    wb = openpyxl.Workbook()

    # ── شیت ۱: خلاصه آمار ────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "خلاصه آمار"
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14

    ws["A1"] = f"📊 گزارش بصیر — جستجو: {query or '—'}"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"تاریخ تهیه گزارش: {_now_jalali()}"
    ws["A2"].font = Font(italic=True, color="666666", size=10)

    summary = [
        ("تعداد کل پست‌ها", s["n"]),
        ("مجموع بازدیدها", _fmt_int(s["total_views"])),
        ("مجموع ری‌اکشن‌ها", _fmt_int(s["total_reactions"])),
        ("ری‌اکشن‌های مثبت 👍", _fmt_int(s["pos"])),
        ("ری‌اکشن‌های منفی 👎", _fmt_int(s["neg"])),
        ("ری‌اکشن‌های خنثی", _fmt_int(s["neu"])),
        ("میانگین بازدید هر پست", _fmt_float(s["avg_views"])),
        ("میانگین ری‌اکشن هر پست", _fmt_float(s["avg_reactions"])),
        ("نرخ تعامل (%)", f"{s['engagement']:.3f}%"),
    ]
    row = 4
    for label, value in summary:
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = _LABEL_FONT
        lc.fill = _LABEL_FILL
        lc.border = _BORDER
        lc.alignment = Alignment(horizontal="right", vertical="center")
        vc = ws.cell(row=row, column=2, value=value)
        vc.border = _BORDER
        vc.alignment = Alignment(horizontal="center", vertical="center")
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="📊 توزیع ری‌اکشن‌ها").font = _TITLE_FONT
    row += 1
    _write_header(ws, ["ایموجی", "تعداد", "نوع"], [28, 22, 14], row=row)
    row += 1
    for emoji, cnt in sorted(s["emoji_counts"].items(), key=lambda x: x[1], reverse=True):
        ws.cell(row=row, column=1, value=emoji).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=2, value=cnt).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=3, value=_reaction_kind(emoji)).alignment = Alignment(horizontal="center")
        for col in (1, 2, 3):
            ws.cell(row=row, column=col).border = _BORDER
        row += 1

    # ── شیت ۲: پست‌ها ─────────────────────────────────────────────────────────
    wp = wb.create_sheet("پست‌ها")
    wp.sheet_view.rightToLeft = True
    headers = ["ردیف", "پلتفرم", "کانال", "تاریخ پست", "متن پست", "بازدید",
               "ری‌اکشن", "ری‌اکشن مثبت", "ری‌اکشن منفی", "فوروارد", "کامنت", "نوع رسانه"]
    widths = [6, 10, 26, 30, 60, 12, 10, 12, 12, 10, 10, 12]
    _write_header(wp, headers, widths)
    for i, r in enumerate(s["rows"], start=1):
        rownum = i + 1
        fill = _ROW_ODD if rownum % 2 == 0 else _ROW_EVEN
        values = [
            i, PLATFORM, r["channel"], _fmt_dt_cell(r["jdate"], r["jtime"]),
            r["text"], r["views"], r["reactions"], r["pos"], r["neg"],
            r["forwards"], r["comments"], r["media"],
        ]
        for col, val in enumerate(values, start=1):
            c = wp.cell(row=rownum, column=col, value=val)
            c.fill = fill
            c.border = _BORDER
            c.alignment = Alignment(
                horizontal="right" if col == 5 else "center",
                vertical="center", wrap_text=True,
            )
        wp.row_dimensions[rownum].height = 38
    wp.freeze_panes = "A2"

    # ── شیت ۳: آمار کانال‌ها ──────────────────────────────────────────────────
    wc = wb.create_sheet("آمار کانال‌ها")
    wc.sheet_view.rightToLeft = True
    _write_header(
        wc,
        ["ردیف", "نام کانال", "تعداد پست", "مجموع بازدید", "مجموع ری‌اکشن", "مجموع فوروارد"],
        [6, 32, 12, 16, 16, 16],
    )
    ch_sorted = sorted(s["per_channel"].items(), key=lambda x: x[1]["views"], reverse=True)
    for i, (name, d) in enumerate(ch_sorted, start=1):
        rownum = i + 1
        fill = _ROW_ODD if rownum % 2 == 0 else _ROW_EVEN
        values = [i, name, d["posts"], d["views"], d["reactions"], d["forwards"]]
        for col, val in enumerate(values, start=1):
            c = wc.cell(row=rownum, column=col, value=val)
            c.fill = fill
            c.border = _BORDER
            c.alignment = Alignment(
                horizontal="right" if col == 2 else "center", vertical="center", wrap_text=True,
            )
    wc.freeze_panes = "A2"

    # ── شیت ۴: توزیع روزانه ───────────────────────────────────────────────────
    wd = wb.create_sheet("توزیع روزانه")
    wd.sheet_view.rightToLeft = True
    _write_header(
        wd, ["تاریخ", "تعداد پست", "مجموع بازدید", "مجموع ری‌اکشن"], [28, 12, 16, 16]
    )
    for i, jdate in enumerate(sorted(s["per_day"].keys()), start=1):
        d = s["per_day"][jdate]
        rownum = i + 1
        fill = _ROW_ODD if rownum % 2 == 0 else _ROW_EVEN
        greg = _greg_of(jdate)
        date_label = f"{jdate} ({greg})" if greg else jdate
        values = [date_label, d["posts"], d["views"], d["reactions"]]
        for col, val in enumerate(values, start=1):
            c = wd.cell(row=rownum, column=col, value=val)
            c.fill = fill
            c.border = _BORDER
            c.alignment = Alignment(
                horizontal="right" if col == 1 else "center", vertical="center",
            )
    wd.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
